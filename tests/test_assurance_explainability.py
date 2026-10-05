import hashlib
import json
import shutil
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from pptx import Presentation

from experian_workflow.assurance import pipeline
from experian_workflow.assurance.explainability import load_explainability


def test_explainability_is_same_run_and_consumed_across_channels():
    root = Path(__file__).resolve().parents[1]
    manifest = pipeline.run_assurance_pipeline(root)
    run_dir = root / "output/assurance_runs" / manifest["run_id"]
    path = run_dir / "publication/explainability.json"
    try:
        guide = load_explainability(run_dir=run_dir, expected_run_id=manifest["run_id"])
        assert guide["code_identity"] == manifest["code_identity"]
        assert guide["as_of_date"] == "2026-09-30"
        assert {s["format"] for s in guide["sources"]} == {"SQLite", "JSON", "CSV"}
        assert guide["metrics"]["sufficient_evidence"]["value"] == 5
        assert guide["metrics"]["sufficient_evidence"]["denominator"] == "all assurance test rows"
        assert guide["metrics"]["overdue_actions"]["trusted_population"] == "remediation_actions"
        assert guide["reconciliation"]["unmapped_population"] == 10
        assert guide["metrics"]["unmapped_tests"]["value"] == 1
        assert "90 days" in guide["states"]["STALE"]
        html = (run_dir / "publication/html/assurance_dashboard.html").read_text()
        assert 'id="data-and-metrics"' in html
        assert "Sources and roles" in html
        assert "Headline metric definitions" in html
        assert "Denominator/context" in html
        assert "10 unmapped source records and 1 unmapped assurance test" in html
        deck = Presentation(run_dir / "publication/powerpoint/assurance_executive_report.pptx")
        assert len(deck.slides) == 8
        slide6 = " ".join(s.text for s in deck.slides[5].shapes if s.has_text_frame)
        assert "SQLite / enterprise reference" in slide6
        assert "Identity normalization" in slide6
        assert "Population reconciliation" in slide6
        assert "Executive claim" in slide6
        assert "One assurance test" in deck.slides[0].notes_slide.notes_text_frame.text
        assert "Denominator/context" in deck.slides[0].notes_slide.notes_text_frame.text
        twb = ET.parse(run_dir / "publication/tableau/assurance_dashboard.twb")
        about = twb.find('.//dashboard[@name="data_trust"]')
        assert about.attrib["caption"] == "About / Data & Metrics"
        sheet = twb.find('.//worksheet[@name="methodology_and_validation"]')
        for field in ("numerator", "denominator", "grain", "how_to_read"):
            assert sheet.find(f'.//column[@name="[{field}]"]') is not None

        with pytest.raises(RuntimeError, match="run identity mismatch"):
            load_explainability(run_dir=run_dir, expected_run_id="stale-run")
        mutated = json.loads(path.read_text())
        del mutated["metrics"]["sufficient_evidence"]["denominator"]
        path.write_text(json.dumps(mutated))
        with pytest.raises(RuntimeError, match="Incomplete explainability metric"):
            load_explainability(run_dir=run_dir, expected_run_id=manifest["run_id"])
        mutated = dict(guide)
        mutated["schema_version"] = "future"
        path.write_text(json.dumps(mutated))
        with pytest.raises(RuntimeError, match="Unsupported or incomplete"):
            load_explainability(run_dir=run_dir, expected_run_id=manifest["run_id"])
        mutated = dict(guide)
        mutated["publication_inputs"] = {}
        mutated["publication_input_identity"] = hashlib.sha256(b"{}").hexdigest()
        path.write_text(json.dumps(mutated))
        with pytest.raises(RuntimeError, match="binding is incomplete"):
            load_explainability(run_dir=run_dir, expected_run_id=manifest["run_id"])
        path.write_text(json.dumps(guide))
        table = run_dir / "publication/tables/assurance_tests.csv"
        table.write_bytes(table.read_bytes() + b"\n")
        with pytest.raises(RuntimeError, match="publication-input bytes changed"):
            load_explainability(run_dir=run_dir, expected_run_id=manifest["run_id"])
    finally:
        shutil.rmtree(run_dir, ignore_errors=True)
