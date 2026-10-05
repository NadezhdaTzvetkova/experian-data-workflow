import json
import re
import shutil
import subprocess
from pathlib import Path

import pytest
from pptx import Presentation

from experian_workflow.assurance.pipeline import run_assurance_pipeline
from experian_workflow.assurance.presentation_validation import (
    validate_presentation_consistency,
)


@pytest.fixture(scope="module")
def publication_run():
    root = Path(__file__).resolve().parents[1]
    manifest = run_assurance_pipeline(root)
    run_dir = root / "output/assurance_runs" / manifest["run_id"]
    yield run_dir
    shutil.rmtree(run_dir)


def test_complete_presentation_parity(publication_run):
    result = validate_presentation_consistency(run_dir=publication_run)
    assert result["status"] == "PASS"
    assert result["headline_metric_count"] == 12
    assert result["html_browser_rendering"] == "NOT_EXECUTED"


def test_javascript_contract_counts_and_distinct_grains(publication_run, tmp_path):
    node = shutil.which("node")
    if node is None:
        pytest.skip("Node is optional for isolated JavaScript arithmetic checks")
    html = (publication_run / "publication/html/assurance_dashboard.html").read_text(encoding="utf-8")

    def payload(element_id):
        return re.search(rf'<script id="{element_id}"[^>]*>(.*?)</script>', html, re.DOTALL).group(1)

    functions = "\n".join(
        re.search(rf"function {name}\(.*?\n\}}", html, re.DOTALL).group(0)
        for name in ("scopedMetrics", "uniqueRecords", "percentage")
    )
    # Execute pure functions only: no page, DOM, browser, chart or interaction emulation.
    script = (
        'const assert = require("node:assert/strict");\n'
        f'const assuranceData = {payload("assurance-data")};\n'
        f'const priorityData = {payload("priority-data")};\n'
        f'const canonicalMetrics = {payload("canonical-metrics")};\n'
        f'const metricDefinitions = {payload("metric-definitions")};\n'
        + functions
        + '\nassert.strictEqual(scopedMetrics(assuranceData), canonicalMetrics);'
        + '\nassert.equal(scopedMetrics([]).assurance_tests, 0);'
        + '\nassert.equal(scopedMetrics(assuranceData.filter(r => r.evidence_sufficiency === "PARTIAL")).partial_evidence, 1);'
        + '\nassert.equal(uniqueRecords(priorityData, "finding_id").length, 5);'
        + '\nassert.equal(uniqueRecords(priorityData, "action_id").length, 6);'
        + '\nassert.equal(percentage(0, 0), "0%");'
    )
    path = tmp_path / "pure-counts.cjs"
    path.write_text(script, encoding="utf-8")
    subprocess.run([node, str(path)], check=True, capture_output=True, text=True)
    # Parse the inline application script without executing it.
    application = re.findall(r"<script>(.*?)</script>", html, re.DOTALL)[-1]
    path.write_text(application, encoding="utf-8")
    subprocess.run([node, "--check", str(path)], check=True, capture_output=True, text=True)


@pytest.mark.parametrize("payload_id", ["canonical-metrics", "assurance-data", "priority-data"])
def test_html_material_mismatch_fails_closed(publication_run, payload_id):
    path = publication_run / "publication/html/assurance_dashboard.html"
    original = path.read_text(encoding="utf-8")
    start = original.index(">", original.index(f'<script id="{payload_id}"')) + 1
    end = original.index("</script>", start)
    payload = json.loads(original[start:end])
    if payload_id == "canonical-metrics":
        payload["assurance_tests"] = 999
    elif payload_id == "assurance-data":
        payload[0]["residual_risk"] = "CRITICAL"
        payload[0]["evaluated_population"] = 999
    else:
        payload.pop()
    try:
        path.write_text(original[:start] + json.dumps(payload) + original[end:], encoding="utf-8")
        with pytest.raises(RuntimeError, match="mismatch"):
            validate_presentation_consistency(run_dir=publication_run)
    finally:
        path.write_text(original, encoding="utf-8")


@pytest.mark.parametrize("target", ["headline", "notes"])
def test_powerpoint_actual_content_mismatch_fails_closed(publication_run, target):
    path = publication_run / "publication/powerpoint/assurance_executive_report.pptx"
    original = path.read_bytes()
    deck = Presentation(path)
    if target == "headline":
        shape = next(s for slide in deck.slides for s in slide.shapes if s.name.startswith("metric:"))
        shape.text = "999"
    else:
        deck.slides[0].notes_slide.notes_text_frame.text = "Stale metric definitions"
    try:
        deck.save(path)
        with pytest.raises(RuntimeError, match="mismatch"):
            validate_presentation_consistency(run_dir=publication_run)
    finally:
        path.write_bytes(original)
