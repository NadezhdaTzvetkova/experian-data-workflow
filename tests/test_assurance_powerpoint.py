import json
import shutil
from pathlib import Path

from experian_workflow.assurance import pipeline
from experian_workflow.assurance.powerpoint import build_powerpoint_publication


def test_powerpoint_publication_uses_governed_current_run_package():
    root = Path(__file__).resolve().parents[1]
    manifest = pipeline.run_assurance_pipeline(root)
    run_dir = root / "output" / "assurance_runs" / str(manifest["run_id"])

    try:
        publication_outputs = manifest["outputs"]["publication"]
        assert "powerpoint_pptx" in publication_outputs
        assert "powerpoint_slide_data" in publication_outputs
        assert "powerpoint_validation" in publication_outputs
        assert (root / publication_outputs["powerpoint_pptx"]).exists()
        assert (root / publication_outputs["powerpoint_slide_data"]).exists()
        assert (root / publication_outputs["powerpoint_validation"]).exists()

        summary = json.loads(
            (run_dir / "publication" / "validation_summary.json").read_text(
                encoding="utf-8"
            )
        )
        assert summary["scope"] == (
            "publication_foundation_html_tableau_and_powerpoint"
        )
        assert summary["powerpoint"]["status"] == "STRUCTURAL_PASS"
        assert summary["powerpoint"]["slide_count"] == 8
        assert summary["powerpoint"]["internal_navigation_links"] == 114
        assert summary["powerpoint"]["run_identity_match"] is True
        assert summary["powerpoint"]["package_structure"] == "PASS"
        assert summary["powerpoint"]["client_validation"] == "PENDING"

        result = build_powerpoint_publication(run_dir=run_dir)
        pptx_path = Path(result["pptx_path"])
        slide_data_path = Path(result["slide_data_path"])

        assert result["status"] == "PASS"
        assert result["slide_count"] == 8
        assert result["internal_navigation_links"] == 114
        assert len(result["external_publication_links"]) == 17
        assert result["run_identity_match"] is True
        assert result["package_structure"] == "PASS"
        assert pptx_path.exists()
        assert slide_data_path.exists()

        slide_data = json.loads(slide_data_path.read_text(encoding="utf-8"))
        metrics = json.loads(
            (run_dir / "publication" / "metrics.json").read_text(
                encoding="utf-8"
            )
        )
        assert slide_data["run_id"] == str(manifest["run_id"])
        assert slide_data["as_of_date"] == metrics["as_of_date"]
        assert slide_data["methodology_version"] == metrics["methodology_version"]
        assert slide_data["metrics"] == metrics["metrics"]
        from pptx import Presentation
        presentation = Presentation(result["pptx_path"])
        slide5_text = " ".join(
            shape.text
            for shape in presentation.slides[4].shapes
            if hasattr(shape, "text")
        )
        assert "Unmapped assurance population: 1" in slide5_text
        assert "Unmapped assurance population: 0" not in slide5_text
        assert slide_data["attention_items"]
        assert all("D:\\" not in link for link in result["external_publication_links"])
    finally:
        shutil.rmtree(run_dir, ignore_errors=True)
