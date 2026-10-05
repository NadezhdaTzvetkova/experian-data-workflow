"""Fail-closed reviewer checks against canonical published records, never test oracles."""

from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from pptx import Presentation


def _payload(html: str, element_id: str) -> object:
    match = re.search(rf'<script id="{re.escape(element_id)}"[^>]*>(.*?)</script>', html, re.DOTALL)
    if match is None:
        raise RuntimeError(f"Missing presentation payload: {element_id}")
    return json.loads(match.group(1))


def validate_presentation_consistency(*, run_dir: Path) -> dict[str, object]:
    publication = run_dir / "publication"
    metrics = json.loads((publication / "metrics.json").read_text(encoding="utf-8"))
    guide = json.loads((publication / "explainability.json").read_text(encoding="utf-8"))
    html = (publication / "html/assurance_dashboard.html").read_text(encoding="utf-8")
    if _payload(html, "canonical-metrics") != metrics["metrics"]:
        raise RuntimeError("HTML canonical headline metric mismatch")
    if _payload(html, "metric-definitions") != guide["metrics"]:
        raise RuntimeError("HTML governed metric definition mismatch")
    for payload_id, table in (
        ("assurance-data", "assurance_tests"),
        ("priority-data", "findings_actions_detail"),
        ("risk-domain-data", "risk_domain_summary"),
    ):
        frame = pd.read_csv(publication / f"tables/{table}.csv")
        expected = json.loads(frame.to_json(orient="records", date_format="iso"))
        if _payload(html, payload_id) != expected:
            raise RuntimeError(f"HTML full-content drill/distribution mismatch: {table}")
    static_cards = {
        "kpi-tests": "assurance_tests", "kpi-sufficient": "sufficient_evidence",
        "kpi-stale": "stale_evidence", "kpi-unmapped": "unmapped_tests",
        "kpi-high-risk": "high_or_critical_residual_risk",
        "kpi-not-evaluable": "not_evaluable_residual_risk",
    }
    for element_id, metric_id in static_cards.items():
        match = re.search(rf'id="{element_id}"[^>]*>(\d+)</div>', html)
        if match is None or int(match.group(1)) != metrics["metrics"][metric_id]:
            raise RuntimeError(f"HTML displayed headline mismatch: {metric_id}")
    accessibility_markers = (
        'class="skip-link"', ':focus-visible', 'aria-label="Finding severity"',
        'role="dialog"', 'aria-modal="true"', 'aria-hidden="true"',
        'onkeydown="activateDrillRow(event, this)"', 'event.key === "Escape"',
        'id="accessible-assurance-table"', 'scope="col"',
    )
    if any(marker not in html for marker in accessibility_markers):
        raise RuntimeError("HTML accessibility structure is incomplete")
    deck = Presentation(publication / "powerpoint/assurance_executive_report.pptx")
    seen = set()
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.name.startswith("metric:"):
                key = shape.name.removeprefix("metric:")
                if shape.text != str(metrics["metrics"][key]):
                    raise RuntimeError(f"PowerPoint displayed headline mismatch: {key}")
                seen.add(key)
        notes = slide.notes_slide.notes_text_frame.text
        for key, item in guide["metrics"].items():
            if f"{item['label']}: {metrics['metrics'][key]} {item['units']}" not in notes:
                raise RuntimeError(f"PowerPoint metric note mismatch: {key}")
    if not set(static_cards.values()).difference({"unmapped_tests", "not_evaluable_residual_risk"}).issubset(seen):
        raise RuntimeError("PowerPoint headline metric bindings are missing")
    return {
        "status": "PASS", "run_id": metrics["run_id"],
        "headline_metric_count": len(metrics["metrics"]),
        "html_complete_records_and_distributions": "PASS",
        "powerpoint_displayed_headlines_and_all_metric_notes": "PASS",
        "html_accessibility_structure": "PASS",
        "html_browser_rendering": "NOT_EXECUTED",
        "tableau_acceptance": "Covered by existing full-content Hyper and workbook validation",
    }
