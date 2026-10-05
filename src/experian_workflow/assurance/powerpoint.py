from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from typing import Any

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from experian_workflow.assurance.explainability import load_explainability

FONT = "Arial"
BG = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(25, 35, 52)
MUTED = RGBColor(92, 108, 130)
BORDER = RGBColor(214, 222, 233)
NAVY = RGBColor(20, 42, 74)
BLUE = RGBColor(36, 99, 235)
BLUE_DARK = RGBColor(30, 64, 175)
BLUE_SOFT = RGBColor(232, 240, 255)
CYAN = RGBColor(14, 116, 144)
CYAN_SOFT = RGBColor(225, 249, 253)
GREEN = RGBColor(21, 128, 61)
GREEN_SOFT = RGBColor(229, 249, 237)
AMBER = RGBColor(180, 83, 9)
AMBER_SOFT = RGBColor(255, 247, 220)
RED = RGBColor(185, 28, 28)
RED_SOFT = RGBColor(255, 234, 234)
PURPLE = RGBColor(126, 34, 206)
PURPLE_SOFT = RGBColor(247, 236, 255)
SLATE_SOFT = RGBColor(239, 243, 248)

SLIDE_TITLES = [
    "Executive Assurance Snapshot",
    "Evidence Quality",
    "Risk & Control Position",
    "Findings & Remediation",
    "Data Trust & Reconciliation",
    "Traceability",
    "Interactive Analytics",
    "Methodology, Limitations & Next Steps",
]

SECTION_TABS = [
    ("Snapshot", 0),
    ("Evidence", 1),
    ("Risk", 2),
    ("Findings", 3),
    ("Trust", 4),
    ("Trace", 5),
    ("Explore", 6),
    ("Methods", 7),
]

LINKS_BY_SLIDE = {
    0: [("HTML dashboard", "../html/assurance_dashboard.html"), ("Tableau", "../tableau/assurance_dashboard.twbx")],
    1: [("Evidence report", "../html/reports/executive_assurance.html"), ("Evidence chart", "../html/charts/evidence_sufficiency.html")],
    2: [("Risk report", "../html/reports/risk_and_evidence.html"), ("Residual risk", "../html/charts/residual_risk.html")],
    3: [("Findings report", "../html/reports/findings_and_actions.html"), ("Actions chart", "../html/charts/remediation_actions.html")],
    4: [("Data-trust report", "../html/reports/data_trust.html"), ("Freshness chart", "../html/charts/freshness.html")],
    5: [("Data-trust report", "../html/reports/data_trust.html")],
    6: [("Open HTML", "../html/assurance_dashboard.html"), ("Open Tableau", "../tableau/assurance_dashboard.twbx")],
    7: [("Open HTML", "../html/assurance_dashboard.html"), ("Open Tableau", "../tableau/assurance_dashboard.twbx")],
}


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _read_table(path: Path) -> list[dict[str, str]]:
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    return [
        {str(key): str(value) for key, value in row.items()}
        for row in frame.to_dict(orient="records")
    ]


def build_slide_data(*, run_dir: Path) -> dict[str, Any]:
    publication_dir = run_dir / "publication"
    tables_dir = publication_dir / "tables"
    metrics = _read_json(publication_dir / "metrics.json")
    metric_validation = _read_json(publication_dir / "metric_validation.json")
    validation_summary = _read_json(publication_dir / "validation_summary.json")
    evidence_summary = _read_json(run_dir / "evidence_summary.json")

    payload = {
        "run_id": str(metrics["run_id"]),
        "as_of_date": str(metrics["as_of_date"]),
        "methodology_version": str(metrics["methodology_version"]),
        "validation_status": str(metrics["validation_status"]),
        "metrics": metrics["metrics"],
        "reconciliation": evidence_summary["reconciliation"],
        "explainability": load_explainability(run_dir=run_dir, expected_run_id=str(metrics["run_id"])),
        "metric_validation_status": str(metric_validation["status"]),
        "publication_validation_status": str(validation_summary["status"]),
        "risk_domain_summary": _read_table(tables_dir / "risk_domain_summary.csv"),
        "attention_items": _read_table(tables_dir / "attention_items.csv"),
        "findings_summary": _read_table(tables_dir / "findings_summary.csv"),
        "remediation_summary": _read_table(tables_dir / "remediation_summary.csv"),
        "lineage_summary": _read_table(tables_dir / "lineage_summary.csv"),
    }

    run_ids: set[str] = set()
    for table_name in (
        "risk_domain_summary",
        "attention_items",
        "findings_summary",
        "remediation_summary",
        "lineage_summary",
    ):
        for row in payload[table_name]:
            value = row.get("publication_run_id", "")
            if value:
                run_ids.add(value)
    if run_ids != {payload["run_id"]}:
        raise RuntimeError(
            f"PowerPoint slide-data run identity mismatch: {sorted(run_ids)}"
        )
    return payload


def _safe_int(value: Any) -> int:
    return int(float(str(value or 0)))


def _rect(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    fill: RGBColor,
    line: RGBColor | None = None,
    radius: bool = False,
) -> Any:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if radius
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line is not None else fill
    return shape


def _text(
    slide: Any,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    size: float = 16,
    bold: bool = False,
    color: RGBColor = INK,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> Any:
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.01)
    frame.margin_right = Inches(0.01)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    box._element.nvSpPr.cNvPr.set("descr", text)
    return box


def _background(slide: Any) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _pill(
    slide: Any,
    text: str,
    left: float,
    top: float,
    width: float,
    *,
    fill: RGBColor,
    color: RGBColor,
) -> Any:
    shape = _rect(slide, left, top, width, 0.28, fill=fill, line=fill, radius=True)
    shape.text_frame.clear()
    p = shape.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(7.5)
    run.font.bold = True
    run.font.color.rgb = color
    return shape


def _button(
    slide: Any,
    text: str,
    left: float,
    top: float,
    width: float,
    *,
    fill: RGBColor = WHITE,
    line: RGBColor = BORDER,
    color: RGBColor = NAVY,
    target_slide: Any | None = None,
    hyperlink: str | None = None,
) -> Any:
    shape = _rect(slide, left, top, width, 0.34, fill=fill, line=line, radius=True)
    shape.text_frame.clear()
    p = shape.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(7.5)
    run.font.bold = True
    run.font.color.rgb = color
    if target_slide is not None:
        shape.click_action.target_slide = target_slide
    elif hyperlink is not None:
        shape.click_action.hyperlink.address = hyperlink
    return shape


def _header(
    slide: Any,
    *,
    title: str,
    subtitle: str,
    slide_no: int,
    run_id: str,
    as_of_date: str,
) -> None:
    _text(slide, title, 0.68, 0.34, 8.7, 0.42, size=23, bold=True, color=NAVY)
    _text(slide, subtitle, 0.70, 0.84, 9.9, 0.25, size=9.5, color=MUTED)
    _pill(slide, f"AS OF {as_of_date}", 10.10, 0.37, 1.55, fill=BLUE_SOFT, color=BLUE_DARK)
    _pill(slide, f"{slide_no:02d}/08", 11.78, 0.37, 0.78, fill=SLATE_SOFT, color=NAVY)
    _rect(slide, 0.70, 1.18, 11.85, 0.025, fill=BLUE)
    _text(slide, f"Run {run_id}", 9.28, 0.86, 3.25, 0.20, size=7.3, color=MUTED, align=PP_ALIGN.RIGHT)


def _section_tabs(prs: Presentation, slide_index: int) -> None:
    slide = prs.slides[slide_index]
    start_x = 0.70
    gap = 0.04
    width = 1.42
    for idx, (label, target_index) in enumerate(SECTION_TABS):
        active = idx == slide_index
        shape = _button(
            slide,
            label,
            start_x + idx * (width + gap),
            1.29,
            width,
            fill=NAVY if active else WHITE,
            line=NAVY if active else BORDER,
            color=WHITE if active else MUTED,
            target_slide=prs.slides[target_index],
        )
        shape.height = Inches(0.30)


def _footer(prs: Presentation, slide_index: int, slide_data: dict[str, Any]) -> None:
    slide = prs.slides[slide_index]
    _rect(slide, 0.68, 6.91, 11.88, 0.018, fill=BORDER)
    prev_index = max(0, slide_index - 1)
    next_index = min(len(prs.slides) - 1, slide_index + 1)
    _button(slide, "◀ Previous", 0.70, 7.01, 1.05, target_slide=prs.slides[prev_index])
    _button(slide, "Home", 1.86, 7.01, 0.70, target_slide=prs.slides[0])
    _button(slide, "Explore", 2.67, 7.01, 0.85, target_slide=prs.slides[6])
    _button(slide, "Next ▶", 3.63, 7.01, 0.95, target_slide=prs.slides[next_index])
    x = 8.65
    for label, target in LINKS_BY_SLIDE[slide_index]:
        _button(slide, label, x, 7.01, 1.55, fill=BLUE_SOFT, line=BLUE_SOFT, color=BLUE_DARK, hyperlink=target)
        x += 1.65
    _text(
        slide,
        f"Validated publication • {slide_data['methodology_version']}",
        4.80,
        7.07,
        3.55,
        0.16,
        size=6.8,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def _kpi_card(
    slide: Any,
    value: Any,
    label: str,
    left: float,
    top: float,
    width: float,
    *,
    accent: RGBColor,
    fill: RGBColor = WHITE,
    target_slide: Any | None = None,
) -> Any:
    card = _rect(slide, left, top, width, 1.02, fill=fill, line=BORDER, radius=True)
    _rect(slide, left, top, 0.055, 1.02, fill=accent, line=accent, radius=True)
    value_shape = _text(slide, str(value), left + 0.18, top + 0.12, width - 0.32, 0.42, size=10.5 if len(str(value)) > 16 else 21, bold=True, color=NAVY)
    _text(slide, label, left + 0.18, top + 0.60, width - 0.32, 0.34, size=10, color=MUTED)
    metric_labels = {
        "Assurance tests": "assurance_tests", "Sufficient evidence": "sufficient_evidence",
        "Stale evidence": "stale_evidence", "High / critical risk": "high_or_critical_residual_risk",
        "Open findings": "open_findings", "Overdue actions": "overdue_actions",
        "Tests using stale evidence": "stale_evidence", "Evidence not evaluable": "not_evaluable_evidence",
        "Repeat findings": "repeat_findings",
    }
    if label in metric_labels:
        value_shape.name = "metric:" + metric_labels[label]
        value_shape._element.nvSpPr.cNvPr.set("descr", f"{label}: {value}")
    if target_slide is not None:
        card.click_action.target_slide = target_slide
    return card


def _callout(
    slide: Any,
    eyebrow: str,
    heading: str,
    body: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    fill: RGBColor,
    accent: RGBColor,
) -> None:
    _rect(slide, left, top, width, height, fill=fill, line=fill, radius=True)
    _pill(slide, eyebrow.upper(), left + 0.16, top + 0.14, min(1.30, width - 0.32), fill=WHITE, color=accent)
    _text(slide, heading, left + 0.16, top + 0.52, width - 0.32, 0.34, size=11.2, bold=True, color=NAVY)
    _text(slide, body, left + 0.16, top + 0.92, width - 0.32, max(0.12, height - 1.04), size=8.4, color=INK)


def _stacked_bar(
    slide: Any,
    segments: list[tuple[str, int, RGBColor]],
    left: float,
    top: float,
    width: float,
    *,
    total: int,
) -> None:
    x = left
    for label, value, color in segments:
        seg_w = width * value / total if total else 0
        if seg_w > 0:
            _rect(slide, x, top, seg_w, 0.42, fill=color, line=color, radius=False)
            if seg_w >= 0.55:
                _text(slide, str(value), x, top + 0.07, seg_w, 0.20, size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
            x += seg_w
    legend_x = left
    for label, value, color in segments:
        _rect(slide, legend_x, top + 0.58, 0.10, 0.10, fill=color, line=color, radius=True)
        _text(slide, f"{label} {value}", legend_x + 0.15, top + 0.52, 1.40, 0.18, size=7.2, color=MUTED)
        legend_x += 1.55


def _bar_row(
    slide: Any,
    label: str,
    value: int,
    max_value: int,
    left: float,
    top: float,
    width: float,
    *,
    fill: RGBColor,
) -> None:
    _text(slide, label, left, top, 1.95, 0.20, size=8.2, color=INK)
    _rect(slide, left + 2.05, top + 0.03, width - 2.55, 0.14, fill=SLATE_SOFT, line=SLATE_SOFT, radius=True)
    bar_width = 0 if max_value <= 0 else (width - 2.55) * value / max_value
    if bar_width > 0:
        _rect(slide, left + 2.05, top + 0.03, bar_width, 0.14, fill=fill, line=fill, radius=True)
    _pill(slide, str(value), left + width - 0.42, top - 0.04, 0.36, fill=WHITE, color=NAVY)


def _table(
    slide: Any,
    rows: list[dict[str, str]],
    columns: list[tuple[str, str]],
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    max_rows: int = 6,
    font_size: float = 7.2,
) -> None:
    display = rows[:max_rows]
    graphic = slide.shapes.add_table(
        len(display) + 1,
        len(columns),
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    table = graphic.table
    for c, (_, label) in enumerate(columns):
        cell = table.cell(0, c)
        cell.text = label
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
    for r, row in enumerate(display, start=1):
        for c, (key, _) in enumerate(columns):
            cell = table.cell(r, c)
            cell.text = row.get(key, "")
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else RGBColor(243, 246, 250)
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.025)
            cell.margin_bottom = Inches(0.025)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font_size if r else font_size + 0.3)
                    run.font.bold = r == 0
                    run.font.color.rgb = WHITE if r == 0 else INK


def _risk_matrix(slide: Any, rows: list[dict[str, str]], left: float, top: float) -> None:
    _text(slide, "DOMAIN SIGNAL MATRIX", left, top, 4.6, 0.20, size=8, bold=True, color=MUTED)
    y = top + 0.34
    headers = ["Domain", "Evidence", "Stale", "High/Crit"]
    widths = [2.35, 0.72, 0.72, 0.80]
    x = left
    for header, width in zip(headers, widths, strict=True):
        _text(slide, header, x, y, width, 0.18, size=7, bold=True, color=MUTED, align=PP_ALIGN.CENTER if x > left else PP_ALIGN.LEFT)
        x += width + 0.08
    y += 0.28
    for row in rows[:6]:
        x = left
        _text(slide, row.get("risk_domain", ""), x, y, widths[0], 0.24, size=7.7, color=INK)
        x += widths[0] + 0.08
        vals = [
            (_safe_int(row.get("sufficient_evidence")), GREEN_SOFT, GREEN),
            (_safe_int(row.get("stale_evidence")), AMBER_SOFT, AMBER),
            (_safe_int(row.get("high_or_critical_residual_risk")), RED_SOFT, RED),
        ]
        for (value, fill, color), width in zip(vals, widths[1:], strict=True):
            _pill(slide, str(value), x, y - 0.03, width, fill=fill, color=color)
            x += width + 0.08
        y += 0.46


def _build_deck(slide_data: dict[str, Any]) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = [prs.slides.add_slide(blank) for _ in SLIDE_TITLES]
    for slide in slides:
        _background(slide)

    m = slide_data["metrics"]
    tests = _safe_int(m["assurance_tests"])
    sufficient = _safe_int(m["sufficient_evidence"])
    partial = _safe_int(m.get("partial_evidence", 0))
    insufficient = _safe_int(m.get("insufficient_evidence", 0))
    not_eval_evidence = _safe_int(m.get("not_evaluable_evidence", 0))
    stale = _safe_int(m["stale_evidence"])
    unmapped = _safe_int(m["unmapped_tests"])
    high_risk = _safe_int(m["high_or_critical_residual_risk"])
    not_eval_risk = _safe_int(m["not_evaluable_residual_risk"])
    open_findings = _safe_int(m["open_findings"])
    repeat_findings = _safe_int(m["repeat_findings"])
    overdue_actions = _safe_int(m["overdue_actions"])

    for idx, title in enumerate(SLIDE_TITLES):
        subtitles = [
            "Decision-ready view of the current validated assurance position",
            "How much of the assurance conclusion is supported by usable, current evidence",
            "Where risk pressure and evidence limitations concentrate by domain",
            "What remains open, recurring or overdue and where management follow-up should focus",
            "Confidence in the data path, mapping state and publication integrity",
            "How executive claims trace back to governed evidence and current-run artifacts",
            "Native PowerPoint navigation into the deeper HTML and Tableau analytical channels",
            "Scope, validation boundaries, limitations and recommended next actions",
        ]
        _header(
            slides[idx],
            title=title,
            subtitle=subtitles[idx],
            slide_no=idx + 1,
            run_id=slide_data["run_id"],
            as_of_date=slide_data["as_of_date"],
        )
        _section_tabs(prs, idx)

    # Slide 1 — Snapshot
    s = slides[0]
    _text(s, "CURRENT POSITION", 0.72, 1.82, 2.1, 0.20, size=8, bold=True, color=MUTED)
    _text(s, "Three signals matter most", 0.72, 2.10, 4.7, 0.38, size=18, bold=True, color=NAVY)
    _callout(s, "Evidence", "Coverage is usable but incomplete", f"{sufficient} of {tests} tests have sufficient evidence. {stale} rely on stale evidence and {not_eval_evidence} remain not evaluable for evidence sufficiency.", 0.72, 2.62, 3.70, 1.68, fill=BLUE_SOFT, accent=BLUE_DARK)
    _callout(s, "Risk", "Residual-risk visibility is constrained", f"{high_risk} tests are high/critical residual risk while {not_eval_risk} cannot yet support a residual-risk conclusion from the available evidence.", 4.63, 2.62, 3.70, 1.68, fill=RED_SOFT, accent=RED)
    _callout(s, "Action", "Follow-up is concentrated", f"The current population contains {open_findings} open findings, {repeat_findings} repeat finding and {overdue_actions} overdue action, enabling targeted rather than broad remediation focus.", 8.54, 2.62, 3.70, 1.68, fill=AMBER_SOFT, accent=AMBER)
    _text(s, "CLICK A KPI TO DRILL INTO THE RELEVANT SECTION", 0.72, 4.62, 5.6, 0.18, size=7.3, bold=True, color=MUTED)
    cards = [
        (tests, "Assurance tests", BLUE, 1),
        (sufficient, "Sufficient evidence", GREEN, 1),
        (stale, "Stale evidence", AMBER, 1),
        (high_risk, "High / critical risk", RED, 2),
        (open_findings, "Open findings", PURPLE, 3),
        (overdue_actions, "Overdue actions", RED, 3),
    ]
    for idx, (value, label, accent, target) in enumerate(cards):
        _kpi_card(s, value, label, 0.72 + idx * 1.96, 4.98, 1.78, accent=accent, target_slide=prs.slides[target])

    _text(s, "Illustrative synthetic data and methodology; not Experian internal.", 0.72, 6.43, 11.53, 0.24, size=9, color=MUTED)

    # Slide 2 — Evidence Quality
    s = slides[1]
    _text(s, "EVIDENCE SUFFICIENCY", 0.72, 1.82, 2.5, 0.20, size=8, bold=True, color=MUTED)
    _text(s, f"{sufficient}/{tests} tests currently have sufficient evidence", 0.72, 2.12, 5.9, 0.34, size=17, bold=True, color=NAVY)
    _stacked_bar(
        s,
        [
            ("Sufficient", sufficient, GREEN),
            ("Partial", partial, AMBER),
            ("Insufficient", insufficient, RED),
            ("Not evaluable", not_eval_evidence, PURPLE),
        ],
        0.72,
        2.65,
        7.25,
        total=max(tests, 1),
    )
    _callout(s, "Interpretation", "Evidence quality is the limiting factor", "The main constraint is not test volume. Stale, insufficient and not-evaluable evidence weakens the confidence of downstream control and residual-risk conclusions.", 8.35, 1.92, 3.90, 1.90, fill=BLUE_SOFT, accent=BLUE_DARK)
    _kpi_card(s, stale, "Tests using stale evidence", 8.35, 4.03, 1.82, accent=AMBER)
    _kpi_card(s, not_eval_evidence, "Evidence not evaluable", 10.36, 4.03, 1.89, accent=PURPLE)
    _callout(s, "Next move", "Repair evidence before broad retesting", "Refreshing stale evidence and resolving not-evaluable cases should improve both assurance confidence and the interpretability of residual risk without unnecessarily expanding scope.", 0.72, 4.58, 7.25, 1.38, fill=WHITE, accent=BLUE_DARK)
    _button(s, "Jump to risk implications ›", 8.35, 5.35, 3.90, fill=NAVY, line=NAVY, color=WHITE, target_slide=prs.slides[2])

    # Slide 3 — Risk & Control Position
    s = slides[2]
    rows = slide_data["risk_domain_summary"]
    max_tests = max([_safe_int(r.get("assurance_tests")) for r in rows] or [1])
    _text(s, "DOMAIN PROFILE", 0.72, 1.82, 2.0, 0.20, size=8, bold=True, color=MUTED)
    _text(s, "Assurance volume and pressure by risk domain", 0.72, 2.10, 5.8, 0.34, size=17, bold=True, color=NAVY)
    for idx, row in enumerate(rows[:6]):
        _bar_row(s, row.get("risk_domain", ""), _safe_int(row.get("assurance_tests")), max_tests, 0.72, 2.70 + idx * 0.48, 5.75, fill=BLUE)
    _risk_matrix(s, rows, 6.90, 2.11)
    pressure = max(rows, key=lambda r: (_safe_int(r.get("stale_evidence")) + _safe_int(r.get("high_or_critical_residual_risk")), _safe_int(r.get("assurance_tests")))) if rows else {}
    _callout(s, "Attention", "Highest displayed attention signal", f"{pressure.get('risk_domain', 'No domain')} has {_safe_int(pressure.get('stale_evidence'))} stale-evidence and {_safe_int(pressure.get('high_or_critical_residual_risk'))} high/critical residual-risk cases in the governed domain summary.", 6.90, 4.88, 5.35, 1.38, fill=RED_SOFT, accent=RED)
    _button(s, "Open attention items ›", 0.72, 5.68, 5.75, fill=NAVY, line=NAVY, color=WHITE, target_slide=prs.slides[6])

    # Slide 4 — Findings & Remediation
    s = slides[3]
    _kpi_card(s, open_findings, "Open findings", 0.72, 1.91, 2.05, accent=PURPLE)
    _kpi_card(s, repeat_findings, "Repeat findings", 2.97, 1.91, 2.05, accent=AMBER)
    _kpi_card(s, overdue_actions, "Overdue actions", 5.22, 1.91, 2.05, accent=RED)
    _callout(s, "Priority", "Persistence matters more than volume", "A repeat finding and an overdue action are escalation signals: they indicate persistence or delayed remediation rather than simply a larger inventory.", 7.56, 1.91, 4.69, 1.55, fill=AMBER_SOFT, accent=AMBER)
    _text(s, "CURRENT MANAGEMENT ACTIONS", 0.72, 3.36, 3.0, 0.20, size=8, bold=True, color=MUTED)
    _table(s, slide_data["remediation_summary"], [("action_id", "Action"), ("finding_id", "Finding"), ("status", "Status"), ("overdue", "Overdue"), ("days_overdue", "Days")], 0.72, 3.70, 7.38, 2.18, max_rows=6, font_size=7.4)
    _callout(s, "Control", "Closure is not just an administrative status", "Where available, closure-validation evidence should remain linked to action status so CLOSED does not become a substitute for evidence that remediation actually worked.", 8.43, 3.70, 3.82, 1.44, fill=WHITE, accent=BLUE_DARK)
    _button(s, "View finding detail ›", 8.43, 5.39, 3.82, fill=NAVY, line=NAVY, color=WHITE, target_slide=prs.slides[6])

    # Slide 5 — Data Trust & Reconciliation
    s = slides[4]
    r = slide_data["reconciliation"]
    _text(s, "Every source record is accounted for", 0.72, 1.88, 11.53, 0.40, size=21, bold=True, color=NAVY)
    stages = [
        ("expected_population", "expected", BLUE_SOFT, BLUE_DARK),
        ("received_population", "received", CYAN_SOFT, CYAN),
        ("mapped_population", "mapped", PURPLE_SOFT, PURPLE),
        ("evaluated_population", "evaluated", GREEN_SOFT, GREEN),
    ]
    for idx, (key, label, fill, accent) in enumerate(stages):
        left = 0.72 + idx * 2.94
        _rect(s, left, 2.56, 2.65, 1.24, fill=fill, line=fill, radius=True)
        _text(s, f"{_safe_int(r[key])} {label}", left + 0.15, 2.90, 2.35, 0.47, size=20, bold=True, color=accent, align=PP_ALIGN.CENTER)
        if idx < 3:
            _text(s, "→", left + 2.66, 2.96, 0.28, 0.35, size=18, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    _rect(s, 0.72, 4.10, 5.55, 1.02, fill=AMBER_SOFT, line=AMBER_SOFT, radius=True)
    _text(s, f"{_safe_int(r['unmapped_population'])} unmapped source records", 0.92, 4.25, 5.15, 0.32, size=16, bold=True, color=NAVY)
    _text(s, "Received = mapped + unmapped source records", 0.92, 4.69, 5.15, 0.22, size=10, color=INK)
    _rect(s, 6.60, 4.10, 5.65, 1.02, fill=AMBER_SOFT, line=AMBER_SOFT, radius=True)
    _text(s, f"{_safe_int(r['not_tested_population'])} not tested (mapped source records)", 6.80, 4.25, 5.25, 0.32, size=16, bold=True, color=NAVY)
    _text(s, "Mapped = evaluated + not tested", 6.80, 4.69, 5.25, 0.22, size=10, color=INK)
    _text(s, f"Separate assurance-test measure: {unmapped} unmapped assurance test. Source-record counts above use a different grain.", 0.72, 5.37, 11.53, 0.34, size=11, color=NAVY)
    _text(s, f"Intermediate gates: {_safe_int(r['structurally_valid_population'])} structurally valid; {_safe_int(r['testable_population'])} testable. Rejected and not-testable populations are retained explicitly.", 0.72, 5.76, 11.53, 0.18, size=9, color=MUTED)
    _rect(s, 0.72, 5.96, 11.53, 0.67, fill=WHITE, line=BORDER, radius=True)
    _text(s, "Trust boundary", 0.92, 6.08, 1.65, 0.24, size=11, bold=True, color=NAVY)
    _text(s, "Run IDs, hashes and lineage establish identity. Independent validation and governed population checks establish analytical correctness.", 2.68, 6.08, 9.33, 0.36, size=10, color=MUTED)

    # Slide 6 — Traceability
    s = slides[5]
    e = slide_data["explainability"]
    _text(s, "Synthetic sources, one governed analytical truth", 0.72, 1.86, 11.53, 0.40, size=20, bold=True, color=NAVY)
    for idx, source in enumerate(e["sources"]):
        left = 0.72 + idx * 2.94
        _rect(s, left, 2.47, 2.65, 1.48, fill=BLUE_SOFT, line=BLUE_SOFT, radius=True)
        _text(s, source["format"] + " / " + source["id"].replace("_", " "), left + 0.14, 2.62, 2.37, 0.45, size=11, bold=True, color=NAVY)
        _text(s, source["role"], left + 0.14, 3.14, 2.37, 0.55, size=10, color=INK)
    _text(s, "PYTHON-DRIVEN FLOW", 0.72, 4.20, 3.5, 0.22, size=10, bold=True, color=MUTED)
    _text(s, " → ".join(e["pipeline_stages"][:4]), 0.72, 4.58, 11.53, 0.36, size=13, bold=True, color=NAVY)
    _text(s, " → ".join(e["pipeline_stages"][4:]), 0.72, 5.06, 11.53, 0.36, size=13, bold=True, color=NAVY)
    _text(s, e["claim_to_evidence"], 0.72, 5.58, 11.53, 0.35, size=11, bold=True, color=NAVY)
    _text(s, "Grain: one test per test ID / period; findings and actions retain their grains.", 0.72, 6.00, 11.53, 0.22, size=10, color=INK)
    _text(s, e["traceability_limit"], 0.72, 6.25, 11.53, 0.23, size=9, color=MUTED)
    _button(s, "Explore current-run evidence ›", 0.72, 6.52, 11.53, fill=NAVY, line=NAVY, color=WHITE, target_slide=prs.slides[6])

    # Slide 7 — Interactive Analytics
    s = slides[6]
    _text(s, "CHOOSE YOUR DRILL-DOWN", 0.72, 1.82, 3.5, 0.20, size=8, bold=True, color=MUTED)
    _text(s, "Use PowerPoint as the executive navigation layer", 0.72, 2.10, 7.2, 0.36, size=17, bold=True, color=NAVY)
    _text(s, "Click a tile to move within the deck, or open the richer browser/Tableau channels for filtering and record-level exploration.", 0.72, 2.55, 8.8, 0.42, size=9, color=MUTED)
    tiles = [
        ("Evidence Quality", "Evidence sufficiency, freshness and confidence", 1, BLUE_SOFT, BLUE_DARK),
        ("Risk & Control", "Domain pressure and residual-risk visibility", 2, RED_SOFT, RED),
        ("Findings", "Open, repeat and overdue remediation signals", 3, AMBER_SOFT, AMBER),
        ("Data Trust", "Lineage, validation and current-run identity", 4, GREEN_SOFT, GREEN),
        ("Traceability", "Claim-to-source evidence chain", 5, PURPLE_SOFT, PURPLE),
        ("Methods", "Scope, limitations and next steps", 7, SLATE_SOFT, NAVY),
    ]
    for idx, (heading, body, target, fill, accent) in enumerate(tiles):
        col = idx % 3
        row = idx // 3
        left = 0.72 + col * 3.86
        top = 3.14 + row * 1.32
        shape = _rect(s, left, top, 3.56, 1.06, fill=fill, line=fill, radius=True)
        shape.click_action.target_slide = prs.slides[target]
        _pill(s, f"{target + 1:02d}", left + 0.16, top + 0.15, 0.42, fill=accent, color=WHITE)
        _text(s, heading, left + 0.70, top + 0.14, 2.55, 0.28, size=10.5, bold=True, color=NAVY)
        _text(s, body, left + 0.70, top + 0.49, 2.55, 0.34, size=7.5, color=MUTED)
    _button(s, "Open interactive HTML dashboard ↗", 0.72, 5.94, 5.57, fill=BLUE, line=BLUE, color=WHITE, hyperlink="../html/assurance_dashboard.html")
    _button(s, "Open Tableau package ↗", 6.69, 5.94, 5.56, fill=NAVY, line=NAVY, color=WHITE, hyperlink="../tableau/assurance_dashboard.twbx")

    # Slide 8 — Methods / Limitations / Next Steps
    s = slides[7]
    _text(s, "METHOD & ACCEPTANCE", 0.72, 1.82, 2.8, 0.20, size=8, bold=True, color=MUTED)
    detail_cards = [
        (slide_data["run_id"], "Publication run", BLUE),
        (slide_data["methodology_version"], "Methodology", PURPLE),
        (slide_data["metric_validation_status"], "Metric validation", GREEN),
        (slide_data["publication_validation_status"], "Publication validation", GREEN),
    ]
    for idx, (value, label, accent) in enumerate(detail_cards):
        _kpi_card(s, value, label, 0.72 + (idx % 2) * 2.72, 2.18 + (idx // 2) * 1.20, 2.45, accent=accent)
    priorities = [
        ("Priority 1 — Restore decision confidence", f"Refresh the {stale} stale-evidence cases and resolve the {not_eval_evidence} evidence-not-evaluable case.", BLUE_SOFT),
        ("Priority 2 — Close risk uncertainty", f"Resolve the {not_eval_risk} residual-risk conclusions that remain not evaluable, starting with areas already carrying high/critical risk.", RED_SOFT),
        ("Priority 3 — Close operational exceptions", f"Address the {overdue_actions} overdue management action and monitor the {unmapped} unmapped assurance test explicitly.", AMBER_SOFT),
    ]
    for idx, (heading, body, fill) in enumerate(priorities):
        top = 2.05 + idx * 1.13
        _rect(s, 6.35, top, 5.90, 1.00, fill=fill, line=fill, radius=True)
        _text(s, heading, 6.53, top + 0.12, 5.54, 0.29, size=12, bold=True, color=NAVY)
        _text(s, body, 6.53, top + 0.48, 5.54, 0.40, size=10, color=INK)
    _text(s, "Governed publication inputs preserve KPI definitions. Native slide navigation links to HTML/Tableau for interactive filtering.", 0.72, 4.74, 5.17, 0.60, size=10, color=MUTED)
    _text(s, "Illustrative synthetic enterprise-assurance dataset and methodology; not Experian internal data or methodology.", 0.72, 6.34, 11.53, 0.34, size=11, color=MUTED)
    _button(s, "Back to executive snapshot", 0.72, 5.72, 2.45, fill=NAVY, line=NAVY, color=WHITE, target_slide=prs.slides[0])
    _button(s, "Open interactive analytics", 3.44, 5.72, 2.55, fill=BLUE_SOFT, line=BLUE_SOFT, color=BLUE_DARK, target_slide=prs.slides[6])

    # Notes expose the complete governed metric guide without crowding executive slides.
    guide = slide_data["explainability"]
    for slide in slides:
        slide.notes_slide.notes_text_frame.text = (
            "HOW TO READ THIS REPORT\n" + guide["disclaimer"] + "\n"
            + "Reporting period: " + guide["reporting_period"] + "\n"
            + "\n".join(
                f"{item['label']}: {item['value']} {item['units']}. {item['grain']}. "
                f"Numerator: {item['numerator']}. Denominator/context: {item['denominator']}. "
                f"{item['limitation']}"
                for item in guide["metrics"].values()
            )
            + "\nSTATE MEANINGS\n" + "\n".join(f"{key}: {value}" for key, value in guide["states"].items())
            + "\nIMPLEMENTATION CHOICES\n" + "\n".join(f"{key}: {value}" for key, value in guide["implementation_choices"].items())
        )
    _text(slides[1], "How to read this: counts are tests; freshness overlaps sufficiency. Full definitions are in Notes and HTML Data & Metrics.", 0.72, 6.48, 11.53, 0.24, size=9, color=MUTED)
    for idx in range(len(slides)):
        _footer(prs, idx, slide_data)
    return prs


def validate_powerpoint_publication(
    *,
    pptx_path: Path,
    slide_data_path: Path,
    expected_run_id: str,
) -> dict[str, Any]:
    if not pptx_path.exists() or pptx_path.stat().st_size == 0:
        raise RuntimeError("PowerPoint publication is missing or empty")
    if not slide_data_path.exists() or slide_data_path.stat().st_size == 0:
        raise RuntimeError("PowerPoint slide_data.json is missing or empty")

    slide_data = _read_json(slide_data_path)
    if str(slide_data["run_id"]) != expected_run_id:
        raise RuntimeError("PowerPoint slide-data run identity mismatch")

    prs = Presentation(pptx_path)
    if len(prs.slides) != len(SLIDE_TITLES):
        raise RuntimeError("Unexpected PowerPoint slide count")

    internal_links = 0
    external_links: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not hasattr(shape, "click_action"):
                continue
            action = shape.click_action
            target_slide = action.target_slide
            if target_slide is not None:
                internal_links += 1
                continue
            address = action.hyperlink.address
            if address:
                external_links.append(address)

    expected_external = [target for links in LINKS_BY_SLIDE.values() for _, target in links] + [
        "../html/assurance_dashboard.html",
        "../tableau/assurance_dashboard.twbx",
    ]
    if sorted(external_links) != sorted(expected_external):
        raise RuntimeError("PowerPoint external publication links differ from contract")
    if any(re.search(r"^[A-Za-z]:[\\/]", link) for link in external_links):
        raise RuntimeError("PowerPoint contains an absolute Windows hyperlink")

    minimum_internal_links = len(SLIDE_TITLES) * (len(SECTION_TABS) + 4)
    if internal_links < minimum_internal_links:
        raise RuntimeError(
            "PowerPoint internal navigation validation failed: "
            f"{internal_links} < {minimum_internal_links}"
        )

    with zipfile.ZipFile(pptx_path, "r") as package:
        members = set(package.namelist())
    required_members = {
        "[Content_Types].xml",
        "ppt/presentation.xml",
        "ppt/slides/slide1.xml",
        f"ppt/slides/slide{len(SLIDE_TITLES)}.xml",
    }
    if not required_members.issubset(members):
        raise RuntimeError("PowerPoint OPC package is structurally incomplete")

    return {
        "status": "PASS",
        "slide_count": len(prs.slides),
        "internal_navigation_links": internal_links,
        "external_publication_links": sorted(external_links),
        "run_identity_match": True,
        "package_structure": "PASS",
    }


def build_powerpoint_publication(*, run_dir: Path) -> dict[str, Any]:
    publication_dir = run_dir / "publication"
    powerpoint_dir = publication_dir / "powerpoint"
    if powerpoint_dir.exists():
        raise RuntimeError(
            f"PowerPoint publication directory already exists: {powerpoint_dir}"
        )
    powerpoint_dir.mkdir(parents=False, exist_ok=False)
    slide_data_path = powerpoint_dir / "slide_data.json"
    pptx_path = powerpoint_dir / "assurance_executive_report.pptx"

    slide_data = build_slide_data(run_dir=run_dir)
    slide_data_path.write_text(
        json.dumps(slide_data, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
        newline="\n",
    )

    prs = _build_deck(slide_data)
    prs.save(pptx_path)

    validation = validate_powerpoint_publication(
        pptx_path=pptx_path,
        slide_data_path=slide_data_path,
        expected_run_id=slide_data["run_id"],
    )
    return {
        **validation,
        "pptx_path": pptx_path,
        "slide_data_path": slide_data_path,
    }
